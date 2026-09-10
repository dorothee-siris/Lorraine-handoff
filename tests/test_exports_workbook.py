# tests/test_exports_workbook.py
r"""
Pass-7a unit pins for `lib.exports.page_workbook` (S-LIB-B,
docs/contract_fragments/lib_api_pass7.md) -- the whole-page xlsx export. openpyxl
round-trip: sheet order (Lecture first), Lecture rows, sheet-name truncation and
uniqueness, and the (bytes, filename) return.

    .venv-pinned\Scripts\python -m pytest tests\test_exports_workbook.py -q

Namespace note: same `_import_streamlit_lib` swap-trick as tests/test_links.py (see its
docstring) -- Streamlit/lib and the repo-root pipeline `lib` package would otherwise
collide under the same `lib` name in `sys.modules`.
"""
from __future__ import annotations

import importlib
import io
import sys
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
STREAMLIT_DIR = ROOT / "Streamlit"


def _import_streamlit_lib(*names: str):
    saved = {k: v for k, v in sys.modules.items() if k == "lib" or k.startswith("lib.")}
    for k in saved:
        del sys.modules[k]
    sys.path.insert(0, str(STREAMLIT_DIR))
    try:
        mods = tuple(importlib.import_module(f"lib.{n}") for n in names)
    finally:
        sys.path.remove(str(STREAMLIT_DIR))
        for k in [k for k in sys.modules if k == "lib" or k.startswith("lib.")]:
            del sys.modules[k]
        sys.modules.update(saved)
    return mods


(exports,) = _import_streamlit_lib("exports")


def test_page_workbook_sheet_order_and_lecture_rows_round_trip():
    lecture = [("Fenêtre", "2019–2023"), ("Instantané", "2026-01-01"), ("URL", "https://example.org")]
    sheets = {"Résumé": pd.DataFrame({"a": [1, 2]}), "Détail": pd.DataFrame({"b": [3]})}

    xlsx_bytes, filename = exports.page_workbook(sheets, lecture, view="zoom")
    assert isinstance(xlsx_bytes, bytes) and xlsx_bytes
    assert filename.endswith(".xlsx")
    assert filename.startswith("lorraine-explorer_zoom_vue_")

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert wb.sheetnames == ["Lecture", "Résumé", "Détail"]

    lecture_ws = wb["Lecture"]
    header = [c.value for c in next(lecture_ws.iter_rows(max_row=1))]
    assert header == ["Clé", "Valeur"]
    rows = [(r[0].value, r[1].value) for r in lecture_ws.iter_rows(min_row=2)]
    assert rows == lecture

    # vacuity: a DIFFERENT lecture list must produce DIFFERENT rows, not the same ones
    other_bytes, _ = exports.page_workbook(sheets, [("X", "Y")], view="zoom")
    other_wb = openpyxl.load_workbook(io.BytesIO(other_bytes))
    other_rows = [(r[0].value, r[1].value) for r in other_wb["Lecture"].iter_rows(min_row=2)]
    assert other_rows != rows
    assert other_rows == [("X", "Y")]


def test_page_workbook_data_sheet_contents():
    sheets = {"Résumé": pd.DataFrame({"a": [1, 2], "b": ["x", "y"]})}
    xlsx_bytes, _ = exports.page_workbook(sheets, [], view="zoom")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Résumé"]
    values = list(ws.iter_rows(values_only=True))
    assert values[0] == ("a", "b")
    assert values[1] == (1, "x")
    assert values[2] == (2, "y")


def test_page_workbook_sheet_name_truncation_and_uniqueness():
    long_a = "x" * 40
    long_b = long_a + "!"  # identical first 31 chars -> forces the disambiguation path
    sheets = {long_a: pd.DataFrame({"v": [1]}), long_b: pd.DataFrame({"v": [2]})}

    xlsx_bytes, _ = exports.page_workbook(sheets, [], view="test")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))

    assert all(len(n) <= 31 for n in wb.sheetnames)
    assert len(set(wb.sheetnames)) == len(wb.sheetnames)  # uniqueness, incl. "Lecture"

    data_sheet_names = wb.sheetnames[1:]  # after "Lecture"
    assert data_sheet_names[0] != data_sheet_names[1]
    # vacuity: two sheets whose names ALREADY differ within 31 chars must be left alone,
    # never needlessly renamed
    distinct_sheets = {"Alpha": pd.DataFrame({"v": [1]}), "Beta": pd.DataFrame({"v": [2]})}
    xlsx_bytes2, _ = exports.page_workbook(distinct_sheets, [], view="test")
    wb2 = openpyxl.load_workbook(io.BytesIO(xlsx_bytes2))
    assert wb2.sheetnames[1:] == ["Alpha", "Beta"]
